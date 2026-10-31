
# ============================================================
# FE5108 Final Project - Part 1
# Generate Final Report and Exhibits
# ============================================================

from pathlib import Path
import pandas as pd

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# 1. Project paths
# ============================================================

ROOT = Path(__file__).resolve().parents[2]
DOC_DIR = ROOT / "part1" / "documents"
TABLE_DIR = ROOT / "output" / "tables"
FIG_DIR = ROOT / "output" / "figures"

DOC_DIR.mkdir(parents=True, exist_ok=True)

REPORT_PATH = DOC_DIR / "Part1_report.docx"
EXHIBIT_PATH = DOC_DIR / "Part1_exhibits.docx"


# ============================================================
# 2. Document formatting
# ============================================================

def configure_document(doc, exhibit=False):
    """Apply the agreed academic report formatting."""

    section = doc.sections[0]

    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9 if exhibit else 11)
    normal.font.color.rgb = RGBColor(35, 35, 35)
    normal.paragraph_format.space_after = Pt(4 if exhibit else 6)
    normal.paragraph_format.line_spacing = 1.1

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(14)
    title.font.bold = True
    title.font.color.rgb = RGBColor(35, 35, 35)

    heading = doc.styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.size = Pt(10 if exhibit else 11)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(35, 35, 35)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)


def shade_cell(cell, fill="E8EDF1"):
    """Apply background shading to a table cell."""

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def format_table(table):
    """Apply compact table formatting."""

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            if row_index == 0:
                shade_cell(cell)

            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8)

                    if row_index == 0:
                        run.bold = True


def add_source_note(doc, text):
    """Add a compact source and methodology note."""

    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_before = Pt(3)

    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(90, 90, 90)


def load_csv(filename, index_col=None):
    """Read previously generated analysis results."""

    path = TABLE_DIR / filename

    if not path.exists():
        raise FileNotFoundError(f"Missing result file: {path}")

    return pd.read_csv(path, index_col=index_col)


def check_figure(filename):
    """Check that a required figure exists."""

    path = FIG_DIR / filename

    if not path.exists():
        raise FileNotFoundError(f"Missing figure: {path}")

    return path


# ============================================================
# 3. Generate Part 1 Report
# ============================================================

report = Document()
configure_document(report, exhibit=False)

report.add_heading("Part 1 Factor Investing", level=0)


def add_section(heading, paragraphs):
    """Add a report section with its analysis."""

    report.add_heading(heading, level=1)

    for paragraph in paragraphs:
        report.add_paragraph(paragraph)


# ============================================================
# Question 1A: Factor Premiums and Factor Decay
# ============================================================

add_section(
    "1A Factor Premiums and Factor Decay",
    [
        (
            "The Fama-French factors show different patterns before and "
            "after 1993. The market premium remains statistically "
            "significant in both periods, while monthly SMB and HML "
            "premiums fall from 0.240% to 0.100% and from 0.435% "
            "to 0.146%, respectively. Neither remains significant "
            "after 1993 (Table 1.1)."
        ),
        (
            "Could greater investor attention explain the decline? "
            "As factor strategies become more widely adopted, demand "
            "for stocks with attractive characteristics may increase. "
            "For example, stronger demand for undervalued stocks may "
            "raise their prices and reduce future expected returns. "
            "However, the results cannot confirm this mechanism or "
            "establish a significant change in mean premiums "
            "between the two periods."
        ),
        (
            "The weaker post-1993 evidence makes SMB and HML less "
            "convincing as standalone investment strategies than "
            "their earlier performance suggested. The market "
            "premium, by contrast, remains statistically significant. "
            "Value-weighted momentum and low-beta decile spreads "
            "are also constructed for Sections 1D and 1E."
        )
    ]
)


# ============================================================
# Question 1B: Factor Index Construction in Practice
# ============================================================

add_section(
    "1B Factor Index Construction in Practice",
    [
        (
            "Academic momentum strategies focus on past winners "
            "and losers. In practice, however, liquidity, trading "
            "costs, and concentration also matter. The MSCI USA "
            "Momentum Index illustrates these trade-offs. Its "
            "universe covers the large- and mid-cap US stocks "
            "in the MSCI USA parent index."
        ),
        (
            "Under the December 2013 methodology, the index uses "
            "6- and 12-month momentum signals, excluding the most "
            "recent month. Each signal is adjusted for local "
            "risk-free rates and scaled by three-year annualized "
            "volatility based on weekly returns. The signals are "
            "standardized, combined equally, and standardized "
            "again. The resulting z-score is winsorized at +/-3 "
            "and converted into a positive Momentum Score. "
            "Stocks with the highest positive unwinsorized "
            "z-scores are selected, subject to a target "
            "constituent count."
        ),
        (
            "Weights depend on both Momentum Scores and "
            "parent-index market capitalization. Consider two "
            "stocks with similar momentum signals: the larger "
            "stock may receive more weight. This can reduce "
            "trading impact by favoring generally more liquid "
            "stocks and improve investment capacity. However, "
            "smaller stocks with stronger momentum may receive "
            "relatively low weights, weakening factor exposure."
        ),
        (
            "The index is normally rebalanced in May and "
            "November, with buffer rules to limit turnover. "
            "An additional rebalance is triggered when the "
            "monthly relative increase in market volatility "
            "exceeds the historical 95th percentile of changes "
            "for the MSCI World reference index. Only the "
            "6-month momentum signal is then used, aiming to "
            "mitigate drawdowns during volatility spikes. "
            "A 5% issuer weight cap limits concentration. "
            "The resulting index therefore differs from "
            "an academic long-short momentum factor."
        ),
        (
            "Source: MSCI Momentum Indexes Methodology "
            "(December 2013), Sections 2-3 and "
            "Appendices II-III."
        )
    ]
)


# ============================================================
# Question 1C: Benefits of Factor Diversification
# ============================================================

add_section(
    "1C Benefits of Factor Diversification",
    [
        (
            "Among the four MSCI factor indexes, Momentum "
            "delivers the highest annualized return (13.38%) "
            "and Sharpe ratio (0.686). Value performs worst, "
            "with the lowest return (9.87%) and Sharpe ratio "
            "(0.505). Why not simply invest in Momentum? "
            "Its position as the strongest performer is "
            "only known after observing the full sample."
        ),
        (
            "An equally weighted portfolio of Value, Minimum "
            "Volatility, Equal Weight, and Momentum provides "
            "an alternative, with monthly rebalancing. "
            "Compared with MSCI USA, its annualized return "
            "increases slightly from 11.22% to 11.37%, "
            "while volatility falls from 15.44% to 14.40%. "
            "This improves the Sharpe ratio from 0.587 "
            "to 0.639 (Table 1.2)."
        ),
        (
            "The portfolio still trails Momentum and Minimum "
            "Volatility on a risk-adjusted basis. Yet it "
            "avoids relying entirely on a single factor. "
            "Given the weaker premiums observed in Section "
            "1A, diversification may be preferable to "
            "selecting a strategy solely by its historical "
            "ranking. The historical results favor "
            "diversification relative to the market "
            "benchmark, although future performance may "
            "differ (Figures 1.1-1.2)."
        )
    ]
)


# ============================================================
# Question 1D: Sensitivity to Factor Portfolio Construction
# ============================================================

add_section(
    "1D Sensitivity to Factor Portfolio Construction",
    [
        (
            "The diversification results raise another "
            "question: do differently named factor indexes "
            "provide different sources of return? Monthly "
            "MSCI factor spreads relative to MSCI USA are "
            "compared with value-weighted academic factor "
            "spreads over January 1995-July 2023."
        ),
        (
            "The correlations generally reflect the intended "
            "strategies. Minimum Volatility and low-beta "
            "have a correlation of 0.815, while Value and "
            "HML reach 0.728. MSCI Momentum is also related "
            "to the academic momentum spread (0.642). "
            "These relationships remain visible despite "
            "differences in portfolio construction "
            "(Table 1.3)."
        ),
        (
            "Equal Weight is less straightforward. Its "
            "correlation with SMB is only 0.322, lower "
            "than with HML (0.431). Although equal "
            "weighting gives smaller companies greater "
            "influence, it does not select stocks directly "
            "by size. Other characteristics may therefore "
            "contribute to its returns."
        ),
        (
            "This matters for the multi-factor portfolio "
            "in Section 1C. Different index labels do not "
            "guarantee independent return sources. The "
            "correlations suggest possible overlap, "
            "although factor regressions would be needed "
            "to measure exposures more directly."
        )
    ]
)


# ============================================================
# Question 1E: Regime Sensitivity of Factor Premia
# ============================================================

add_section(
    "1E Regime Sensitivity of Factor Premia",
    [
        (
            "Momentum's historical performance may conceal "
            "periods of substantial weakness. Daniel and "
            "Moskowitz (2016) document momentum crashes "
            "during sharp market rebounds, when previous "
            "losers can recover faster than previous "
            "winners. The 2009 rebound is one example, "
            "although market reversals are not directly "
            "tested here."
        ),
        (
            "Momentum earns an average monthly return "
            "of 1.86% in medium-volatility periods but "
            "-1.03% in high-volatility periods. The "
            "low-volatility group earns only 1.04%, "
            "so the pattern is not strictly monotonic "
            "(Table 1.4). The high-minus-low difference "
            "is not statistically significant under "
            "the Welch test (p = 0.103)."
        ),
        (
            "Using 332 months from December 1995 to July "
            "2023, a regression on trailing 12-month "
            "annualized market volatility produces a "
            "negative coefficient (-0.267). It remains "
            "significant using Newey-West HAC standard "
            "errors with 12 lags (p = 0.022), although "
            "R-squared is only 3.15%. The regression "
            "suggests a negative linear association, "
            "while the regime comparison is less "
            "conclusive."
        ),
        (
            "Higher volatility may therefore warrant "
            "closer monitoring of momentum risk. However, "
            "volatility alone explains little of the "
            "variation in returns. Unlike Daniel and "
            "Moskowitz (2016), this analysis does not "
            "directly identify market reversals or "
            "momentum crashes. Since the volatility "
            "measure includes the current month, "
            "predictive value remains untested "
            "(Figure 1.3)."
        ),
        (
            "Reference: Daniel, K., and Moskowitz, T. J. "
            "(2016). Momentum Crashes. Journal of "
            "Financial Economics, 122(2), 221-247."
        )
    ]
)

report.save(REPORT_PATH)


# ============================================================
# 4. Generate Part 1 Exhibits
# ============================================================

exhibits = Document()
configure_document(exhibits, exhibit=True)

exhibits.add_heading("Part 1 Selected Exhibits", level=0)


# ============================================================
# Table 1.1: Factor Premiums
# ============================================================

results = load_csv("tab_1A_factor_premiums.csv")

required_columns = {
    "Period", "Factor", "N",
    "Mean Monthly Return (%)", "t-stat"
}

if not required_columns.issubset(results.columns):
    raise ValueError("Missing columns in 1A factor results.")

if len(results) != 6:
    raise ValueError("Expected six factor-period observations.")

exhibits.add_heading(
    "Table 1.1 Factor premiums before and after publication",
    level=1
)

table_1A = exhibits.add_table(rows=1, cols=5)

headers = [
    "Factor", "Period", "N",
    "Mean monthly return", "t-stat"
]

for i, header in enumerate(headers):
    table_1A.rows[0].cells[i].text = header

period_names = {
    "1A_pre": "1927–1992",
    "1A_post": "1993–Jul 2023"
}

for period in ["1A_pre", "1A_post"]:
    for factor in ["Mkt-RF", "SMB", "HML"]:

        selected = results.loc[
            (results["Period"] == period) &
            (results["Factor"] == factor)
        ]

        if len(selected) != 1:
            raise ValueError(
                f"Missing or duplicate result: {period}, {factor}"
            )

        record = selected.iloc[0]
        t_stat = record["t-stat"]
        significance = "*" if abs(t_stat) > 1.96 else ""

        values = [
            factor,
            period_names[period],
            str(int(record["N"])),
            f'{record["Mean Monthly Return (%)"]:.3f}%',
            f"{t_stat:.3f}{significance}"
        ]

        cells = table_1A.add_row().cells

        for i, value in enumerate(values):
            cells[i].text = value

format_table(table_1A)

add_source_note(
    exhibits,
    "Source and scope: Kenneth R. French Data Library; "
    "US monthly factor returns, Jan 1927–Jul 2023. "
    "Arithmetic mean monthly returns; t-statistics "
    "without autocorrelation adjustment. "
    "* = significant at 5% (two-sided, approximate "
    "critical value |t| > 1.96)."
)


# ============================================================
# Table 1.2: MSCI Strategy Performance
# ============================================================

performance_1C = load_csv("tab_1C_performance.csv")

fig_return = check_figure("fig_1C_return.png")
fig_sharpe = check_figure("fig_1C_sharpe.png")

exhibits.add_heading(
    "Table 1.2 Annualized performance of MSCI factor indexes",
    level=1
)

table_1C = exhibits.add_table(rows=1, cols=4)

headers_1C = [
    "Strategy",
    "Annualized return",
    "Annualized volatility",
    "Sharpe ratio"
]

for i, header in enumerate(headers_1C):
    table_1C.rows[0].cells[i].text = header

strategy_names = {
    "Market": "MSCI USA",
    "Value": "Value",
    "MinVol": "Minimum volatility",
    "EqualWeight": "Equal weight",
    "Momentum": "Momentum",
    "MultiFactor": "Multi-factor",
    "RiskFree": "Risk-free"
}

strategy_order = [
    "Market",
    "Value",
    "MinVol",
    "EqualWeight",
    "Momentum",
    "MultiFactor",
    "RiskFree"
]

for strategy in strategy_order:

    selected = performance_1C.loc[
        performance_1C["Strategy"] == strategy
    ]

    if len(selected) != 1:
        raise ValueError(f"Missing strategy: {strategy}")

    record = selected.iloc[0]
    sharpe = record["Sharpe Ratio"]

    values = [
        strategy_names[strategy],
        f'{record["Annualized Return (%)"]:.2f}%',
        f'{record["Annualized Volatility (%)"]:.2f}%',
        f"{sharpe:.3f}" if pd.notna(sharpe) else "—"
    ]

    cells = table_1C.add_row().cells

    for i, value in enumerate(values):
        cells[i].text = value

format_table(table_1C)

add_source_note(
    exhibits,
    "Source and scope: MSCI and Kenneth French Data Library. "
    "MSCI USA large- and mid-cap equity universe; "
    "MSCI USA and four factor indexes, gross total returns "
    "in USD, Jan 1995–Jul 2023 (343 months). "
    "Multi-factor: equal-weighted across four factor indexes, "
    "rebalanced monthly. Annualized arithmetic returns = "
    "monthly mean × 12; annualized volatility = monthly "
    "sample standard deviation × sqrt(12). "
    "Sharpe ratios use the Fama-French risk-free rate."
)


# ============================================================
# Figures 1.1 and 1.2
# ============================================================

exhibits.add_heading(
    "Figure 1.1 Annualized arithmetic returns by strategy",
    level=1
)

exhibits.add_picture(
    str(fig_return),
    width=Inches(6.0)
)

exhibits.add_heading(
    "Figure 1.2 Risk-adjusted performance by strategy",
    level=1
)

exhibits.add_picture(
    str(fig_sharpe),
    width=Inches(6.0)
)


# ============================================================
# Table 1.3: Factor Correlations
# ============================================================

corr_1D = load_csv(
    "tab_1D_correlations.csv",
    index_col=0
)

expected_rows = [
    "Value", "MinVol", "EqualWeight", "Momentum"
]

expected_cols = [
    "HML", "SMB", "Momentum", "LowBeta"
]

corr_1D = corr_1D.loc[expected_rows, expected_cols]

assert corr_1D.shape == (4, 4)
assert corr_1D.notna().all().all()

exhibits.add_heading(
    "Table 1.3 Correlation between MSCI and academic factor spreads",
    level=1
)

table_1D = exhibits.add_table(rows=1, cols=5)

headers_1D = [
    "MSCI spread",
    "HML",
    "SMB",
    "Momentum",
    "Low-beta"
]

for i, header in enumerate(headers_1D):
    table_1D.rows[0].cells[i].text = header

row_names = {
    "Value": "Value",
    "MinVol": "Minimum volatility",
    "EqualWeight": "Equal weight",
    "Momentum": "Momentum"
}

for strategy in expected_rows:

    values = [
        row_names[strategy],
        *[
            f"{corr_1D.loc[strategy, col]:.3f}"
            for col in expected_cols
        ]
    ]

    cells = table_1D.add_row().cells

    for i, value in enumerate(values):
        cells[i].text = value

format_table(table_1D)

# Highlight the intended factor pairs in green.
for row_num, col_num in [
    (1, 1),
    (2, 4),
    (3, 2),
    (4, 3)
]:
    shade_cell(
        table_1D.cell(row_num, col_num),
        fill="D9EEE8"
    )

add_source_note(
    exhibits,
    "Source and scope: MSCI and Kenneth French Data Library. "
    "MSCI USA large- and mid-cap equity universe; "
    "monthly US factor return spreads, Jan 1995–Jul 2023 "
    "(343 months). MSCI spreads are factor index returns "
    "minus MSCI USA returns. Academic spreads are HML, SMB, "
    "high-minus-low momentum and low-minus-high beta. "
    "Pearson correlations; corresponding factor pairs "
    "are highlighted in green."
)


# ============================================================
# Table 1.4: Momentum by Volatility Regime
# ============================================================

regime_1E = load_csv(
    "tab_1E_regime_summary.csv",
    index_col=0
)

regression_1E = load_csv(
    "tab_1E_regression.csv",
    index_col=0
)

fig_regime = check_figure("fig_1E_regime_momentum.png")

expected_regimes = [
    "Low",
    "Medium",
    "High"
]

regime_1E = regime_1E.loc[expected_regimes]

assert regime_1E["N"].sum() == 332
assert regime_1E.notna().all().all()

exhibits.add_heading(
    "Table 1.4 Momentum performance by market volatility regime",
    level=1
)

table_1E = exhibits.add_table(rows=1, cols=4)

headers_1E = [
    "Regime",
    "Months",
    "Mean market vol.",
    "Mean momentum"
]

for i, header in enumerate(headers_1E):
    table_1E.rows[0].cells[i].text = header

for regime in expected_regimes:

    record = regime_1E.loc[regime]

    values = [
        regime,
        str(int(record["N"])),
        f'{record["Mean Market Vol (%)"]:.2f}%',
        f'{record["Mean Momentum (%)"]:+.2f}%'
    ]

    cells = table_1E.add_row().cells

    for i, value in enumerate(values):
        cells[i].text = value

format_table(table_1E)

add_source_note(
    exhibits,
    "Source and scope: Kenneth French Data Library, "
    "US value-weighted momentum decile portfolios, "
    "Dec 1995–Jul 2023 (332 months). Market volatility "
    "is the annualized standard deviation of trailing "
    "12 monthly Mkt-RF observations. Regimes are "
    "sample terciles. Momentum is the value-weighted "
    "highest-minus-lowest decile return spread. "
    "High-minus-low Welch test: p = 0.103."
)

ols_row = regression_1E.loc["Conventional OLS"]
hac_row = regression_1E.loc["Newey-West HAC"]

exhibits.add_paragraph(
    "Regression: beta = "
    f'{hac_row["Beta"]:.3f}; '
    f'OLS p = {ols_row["p-value"]:.3f}; '
    f'HAC(12) t = {hac_row["t-stat"]:.3f}, '
    f'p = {hac_row["p-value"]:.3f}; '
    f'R-squared = {hac_row["R-squared"]:.2%}.'
)


# ============================================================
# Figure 1.3: Momentum Returns by Volatility Regime
# ============================================================

exhibits.add_heading(
    "Figure 1.3 Momentum returns across volatility regimes",
    level=1
)

exhibits.add_picture(
    str(fig_regime),
    width=Inches(6.0)
)

exhibits.save(EXHIBIT_PATH)


# ============================================================
# 5. Final Verification
# ============================================================

report_check = Document(REPORT_PATH)
exhibit_check = Document(EXHIBIT_PATH)

report_text = "\n".join(
    p.text for p in report_check.paragraphs
)

checks = {
    "1A investment conclusion":
        "less convincing as standalone" in report_text,

    "1B standardization":
        "standardized again" in report_text,

    "1B volatility trigger":
        "95th percentile" in report_text,

    "1B MSCI World reference":
        "MSCI World reference index" in report_text,

    "1B investment capacity":
        "investment capacity" in report_text,

    "1C best and worst factors":
        "lowest return (9.87%)" in report_text,

    "1C final interpretation":
        "future performance may differ" in report_text,

    "1D correlations":
        "0.815" in report_text,

    "1E statistical analysis":
        "Newey-West" in report_text,

    "1E literature reference":
        "Daniel and Moskowitz" in report_text,

    "Four exhibit tables":
        len(exhibit_check.tables) == 4,

    "Three exhibit figures":
        len(exhibit_check.inline_shapes) == 3,
}

print("\nFINAL VERIFICATION")
print("-" * 45)

for name, passed in checks.items():
    status = "PASS" if passed else "FAIL"
    print(f"{status}: {name}")

assert all(checks.values()), (
    "Document verification failed. Review the FAIL items."
)

print("\nAll final checks passed.")
print("\nGenerated documents:")
print(REPORT_PATH)
print(EXHIBIT_PATH)
print("\nNote: Check the final page count in Microsoft Word.")
