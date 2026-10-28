"""Consolidate saved FE5108 Part 2 outputs; no analysis scripts are rerun.
Run from any directory with the established bundled Python environment.
"""
from pathlib import Path
import hashlib,html,json,math,re,shutil,sys
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'part2/generated/consolidated'
DOC_DIR=ROOT/'part2/documents'
sys.path.insert(0,str(ROOT))
from src import config,utils
TAB=config.TAB_DIR
CHECKS={};LEDGER=[]


def check(name,cond):
    if not bool(cond):raise AssertionError(name)
    CHECKS[name]='PASS'


def read(name,**kwargs):return pd.read_csv(TAB/name,**kwargs)


def number(label,value,fmt,source):
    text=format(value,fmt);LEDGER.append(dict(claim=label,value=float(value),display=text,source=source));return text


def style_doc(fontsize=11):
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.2677);s.page_height=Inches(11.6929)
    s.top_margin=s.bottom_margin=Inches(.85);s.left_margin=s.right_margin=Inches(1)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    d.styles['Normal'].font.size=Pt(fontsize)
    d.styles['Normal'].paragraph_format.line_spacing=1.08
    d.styles['Normal'].paragraph_format.space_after=Pt(6)
    d.styles['Title'].font.size=Pt(15);d.styles['Title'].paragraph_format.space_after=Pt(9)
    d.styles['Heading 1'].font.size=Pt(11);d.styles['Heading 1'].paragraph_format.space_before=Pt(8);d.styles['Heading 1'].paragraph_format.space_after=Pt(4)
    for border in list(d.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    d.core_properties.title='FE5108 Part 2 Security Level Investing'
    d.core_properties.author='FE5108 Part 2'
    return d


def note(d,text):
    p=d.add_paragraph(text);p.paragraph_format.space_after=Pt(5)
    for r in p.runs:r.font.size=Pt(6.5);r.font.color.rgb=RGBColor.from_string('555555')


def addtable(d,frame,widths):
    t=d.add_table(rows=1,cols=len(frame.columns));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for i,(cell,name) in enumerate(zip(t.rows[0].cells,frame.columns)):cell.text=str(name);cell.width=Inches(widths[i])
    hdr=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hdr)
    for row in frame.itertuples(index=False,name=None):
        for i,(cell,value) in enumerate(zip(t.add_row().cells,row)):
            cell.text=str(value);cell.width=Inches(widths[i])
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margins.append(e)
            tcPr.append(margins)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF2');tcPr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.04;p.paragraph_format.keep_with_next=True
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(8);run.bold=ri==0
    spacer=d.add_paragraph();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=Pt(2);spacer.paragraph_format.keep_with_next=True
    spacer.add_run().font.size=Pt(2)
    return t


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    DOC_DIR.mkdir(parents=True,exist_ok=True)
    qA=json.loads((ROOT/'part2/validation_2A.json').read_text())
    qB=json.loads((ROOT/'part2/validation_2B.json').read_text())
    qa=read('tab_2A_monthly_validation.csv',index_col='month')
    a5=read('tab_2A5_summary.csv');a5p=a5[a5.test.eq('P/B vs log(cap rank / book rank)')].set_index('sample')
    conditional=read('tab_2A_pb_conditional_low_bv.csv')
    bp=read('tab_2B_performance.csv',index_col='series')
    cp=read('tab_2C_performance_main.csv',index_col='strategy')
    robust=read('tab_2C_performance_all_cases.csv')
    final=read('tab_2C_final_weights_summary.csv');final=final[final.universe.eq('All observations (main)')].set_index('strategy')
    dp=read('tab_2D_tracking_error_and_hml.csv').set_index(['return_version','strategy'])
    corr=read('tab_2D_correlation_main.csv',index_col=0)
    c=read('tab_2C_monthly_returns_main.csv',index_col='month');b=read('tab_2B_monthly_returns.csv',index_col='month')
    a=read('tab_2D_active_and_hml_main.csv',index_col='month')
    target=pd.period_range('1999-01','2018-12',freq='M')
    for label,x in [('2B',b),('2C',c),('2D',a)]:
        check(label+'_same_240_complete_earned_months',len(x)==240 and pd.PeriodIndex(x.index,freq='M').equals(target) and np.isfinite(x.to_numpy()).all())
    strategies=['EW top 20%','Value score','Value score x cap'];market='2B broad market'
    labels=['S1: EW top 20%','S2: Value score','S3: Value x cap']
    for s,l in zip(strategies,labels):
        check('active_identity_'+s,np.allclose(a[l],c[s]-b['Market: winsorized (main)'],atol=1e-12,rtol=0))
        check('TE_reconciles_'+s,np.isclose(utils.annualized_volatility(a[l]),dp.loc[('Winsorized (main)',s),'annualized_tracking_error'],atol=1e-12))
        check('arithmetic_CAGR_volatility_reconcile_'+s,np.allclose([utils.annualized_return(c[s],method='arithmetic'),utils.annualized_return(c[s],method='geometric'),utils.annualized_volatility(c[s])],cp.loc[s,['annualized_arithmetic_return','annualized_geometric_return','annualized_volatility']].to_numpy(dtype=float),atol=1e-10))
    check('market_identical_in_2B_2C',np.allclose(c[market],b['Market: winsorized (main)'],atol=1e-12))
    check('complete_4_by_4_matrix_reconciles',corr.shape==(4,4) and np.allclose(corr,a.corr(),atol=1e-11) and np.allclose(corr,corr.T) and np.allclose(np.diag(corr),1))
    check('main_final_holdings_222_1109_1109',final.loc[strategies].n_holdings.tolist()==[222,1109,1109])
    check('main_final_weights_nonnegative_unit_sum',final.min_weight.ge(0).all() and final.weight_sum.eq(1).all())
    check('2A_observed_return_coverage_100_percent',qa.coverage_present.eq(1).all() and qa.missing_return_rows.eq(0).all())
    check('2A_full_history_and_2B_window_flags_distinct',qA['rows']==283380 and qB['selected_flagged_stock_months']==5630)
    check('2A5_exact_provisional_flag_comparison_not_all_midpoints',np.isclose(a5p.loc['Without provisional flags','median_rho'],.850339748,atol=1e-9,rtol=0))
    gaps=read('tab_2A_interior_gaps.csv')
    rawstats=read('tab_2A_return_summary.csv',index_col='series')
    flagstats=read('tab_2A_pb_provisional_flag_summary.csv')
    check('2A_counts_ranges_gaps_and_flags_match_prose',len(gaps)==1514 and gaps.stock_id.nunique()==205 and rawstats.loc['Raw','min']==-.922 and rawstats.loc['Raw','max']==30.176 and flagstats.n_flagged.iloc[0]==5720 and qA['lower_clipped']+qA['upper_clipped']==5723)
    check('2A_conditional_frequencies_match_prose',format(conditional.loc[conditional.Bv.eq(.01),'frequency_pb_eq_050_given_bv'].iloc[0],'.2%')=='99.59%' and format(conditional.loc[conditional.Bv.eq(.04),'frequency_pb_eq_050_given_bv'].iloc[0],'.2%')=='6.51%')
    rawcase=robust[robust.universe.eq('All observations (main)')&robust.return_version.eq('Original (sensitivity)')].set_index('strategy')
    flagcase=robust[robust.universe.eq('Without provisional flags (sensitivity)')&robust.return_version.eq('Winsorized (main)')].set_index('strategy')
    check('raw_and_flag_sensitivity_prose_reconciles',np.array_equal(np.round(rawcase.loc[strategies,'annualized_arithmetic_return']*100,2),[24.85,17.24,12.12]) and np.array_equal(np.round(flagcase.loc[strategies,'annualized_arithmetic_return']*100,2),[18.42,14.14,11.01]))
    check('final_maximum_weights_reconcile_to_prose',[format(final.loc[x,'max_weight'],'.4%') for x in strategies]==['0.4505%','0.1763%','0.3761%'])
    msci=read('tab_2B_msci_reference_values.csv',index_col='month');hml=read('tab_2D_hml_reference_values.csv',index_col='month')
    reference=msci.rename(columns={'decimal_return':'MSCI_USA_decimal'}).join(hml)
    reference['MSCI_USA_percent']=reference.MSCI_USA_decimal*100
    check('shared_MSCI_reference_level_ratios',np.allclose(reference.current_level/reference.prior_level-1,reference.MSCI_USA_decimal,atol=1e-12))
    check('shared_HML_percent_decimal_reconcile',np.allclose(reference.HML_percent/100,reference.HML_decimal,atol=1e-12))
    reference.to_csv(OUT/'member1_reference_values.csv',float_format='%.12g')
    # Format report values from their saved owners and retain a claim ledger.
    def C(s,col,fmt='.2%'):return number(s+' '+col,cp.loc[s,col],fmt,'tab_2C_performance_main.csv')
    def B(s,col,fmt='.2%'):return number(s+' '+col,bp.loc[s,col],fmt,'tab_2B_performance.csv')
    def D(s,col,fmt='.2%'):return number(s+' '+col,dp.loc[('Winsorized (main)',s),col],fmt,'tab_2D_tracking_error_and_hml.csv')
    P=[]
    P.append(('Scope', 'We study how data cleaning and portfolio design affect a Value strategy. The backtests use 240 investment months from January 1999 to December 2018, formed one month earlier. Arithmetic annualized return compares average monthly performance; geometric CAGR describes compounded investor growth. All main portfolios use the established winsorized returns and retain the provisional P/B flags.'))
    P.append(('2A Data quality', 'The full extract has 283,380 stock-months for 1,207 securities over 245 formation months, November 1998-March 2019. Monthly membership ranges from 552 to 1,196. April 1999 combines 182 departures with 166 arrivals; the departing IDs subsequently reappear. Counts fall sharply in December 2018 and February 2019, outside the later formation window. These patterns question historical universe consistency. Return coverage is 100% among supplied rows, but 1,514 interior stock-month rows are absent across 205 securities. Portfolios use only supplied formation-month rows and normalize the strategy weights within that observed cross-section; next-month presence is not required. Absent rows are not inserted. Zero-return imputation would assume unchanged investment value for an unobserved outcome, and disappearance does not identify delisting. Dependence on supplied membership still leaves selection and survivorship bias unresolved. Raw returns span -92.2% to +3,017.6%. Monthly linear-interpolated 1st/99th-percentile clipping changes 5,723 observations (2.02%), preserves every row, and uses ex-post cutoffs. This retains the central 98%; a central-99% alternative is documented. See T1 and F1-F3.'))
    P.append(('2A Valuation interpretation', 'P/B has an unusual mass at percentile 0.50: 10,197 observations (3.60%). Conditional midpoint frequency is 99.59% at book-value rank 0.01 but 6.51% at 0.04. The provisional rule Pb=0.50 and Bv<=0.03 flags 5,720 observations, including 5,630 in the investment sample. It may reflect special coding or nonpositive equity, but percentile ranks reveal neither the cause nor equity signs. Median monthly Spearman correlation between P/B and log(cap-score/book-score) is 0.829, rising to 0.850 without these exact flags. Conditional cap/book-rank checks have the expected signs. These are ordinal consistency checks, not verification of P/B = Market Cap / Book Equity: nonlinear ranks do not preserve ratios, and cap scores use a three-month average. See T2 and F4-F5.'))
    P.append(('2B Broad market', f'Normalizing market-cap scores creates a fully invested market proxy. Its correlation with MSCI USA gross USD is {B("Market: winsorized (main)","monthly_correlation_to_msci",".4f")}. Arithmetic annual return is {B("Market: winsorized (main)","annualized_arithmetic_return")} versus {B("MSCI USA gross USD","annualized_arithmetic_return")}; tracking error is {B("Market: winsorized (main)","tracking_error_to_msci")}. An initial $1 becomes ${B("Market: winsorized (main)","terminal_value_of_1",".2f")} versus ${B("MSCI USA gross USD","terminal_value_of_1",".2f")}. This supports aggregate US-market-like co-movement, while percentile weights compress size differences and cannot reconstruct actual capitalization weights. Anonymized selection and unknown corporate-event treatment prevent proving complete coverage or individual-return accuracy. See T3 and F6.'))
    ar=', '.join(C(s,'annualized_arithmetic_return') for s in strategies);vol=', '.join(C(s,'annualized_volatility') for s in strategies)
    P.append(('2C Value construction', f'We reverse the supplied grid using Value score v=1.01-Pb: the cheapest bucket receives 1.00 and the most expensive 0.01, preserving ties. This does not take reciprocals of percentile scores. The three portfolios equally weight ceil(20% of N) cheapest stocks, weight proportionally to v, or weight proportionally to v times the cap score. Cutoff ties in the first portfolio are resolved by stock ID and audited separately. The concentrated cheap-stock selection has the highest mean return and volatility; cap modulation moves the observed score mix closer to the market. See T3 and F7.'))
    P.append(('2C Concentration and robustness', 'At November 2018 formation, EW top 20% has both the fewest holdings and the largest single-stock weight: 222 and 0.4505%. Value-score and Value-times-cap each hold 1,109 stocks, with maxima of 0.1763% and 0.3761%. This is expected here: concentrating on one fifth of the universe assigns about 1/222 to each selected stock, while both score methods spread weights across the full universe. The maximum-weight ordering need not hold in every sample. Lower cap-rank exposure does not establish a causal small-cap premium. Original-return sensitivity raises performance substantially, especially for EW top 20%; provisional-flag exclusions preserve the strategy ordering. Neither result validates the flag or proves a Value premium free of selection bias and costs. See T4-T5.'))
    tes=', '.join(D(s,'annualized_tracking_error') for s in strategies);hs=', '.join(D(s,'HML_correlation','.3f') for s in strategies)
    P.append(('2D Active returns', f'Value-minus-market annualized tracking errors are {tes}. Value times cap has the lowest deviation, plausibly because it retains broad holdings and uses the market cap score in its tilt. Pairwise active-return correlations range from 0.759 to 0.958; correlations with HML are {hs}. The cap-modulated strategy is closest to HML, consistent with value weighting within its component portfolios. HML, however, averages small and big high-minus-low book-to-market portfolios, uses positive book equity, and is long-short; our strategies are monthly-rebalanced long-only portfolios measured against a rank-weighted market. A common subtracted benchmark also contributes to active co-movement. Correlations support a shared Value-related component, not identical exposures or causation. Original-return sensitivity preserves the rankings. Independent reconciliation of MSCI USA and HML with Member 1 remains pending. See T6-T7 and F8.'))
    def refs(text):
        text=re.sub(r'\bT(\d)([ab]?)-T(\d)([ab]?)\b',lambda m:f'Tables 2.{m[1]}{m[2]}-2.{m[3]}{m[4]}',text)
        text=re.sub(r'\bF(\d)-F(\d)\b',lambda m:f'Figures 2.{m[1]}-2.{m[2]}',text)
        text=re.sub(r'\bT2\b','Tables 2.2a-2.2b',text)
        text=re.sub(r'\bT(\d)([ab]?)\b',lambda m:f'Table 2.{m[1]}{m[2]}',text)
        return re.sub(r'\bF(\d)\b',lambda m:f'Figure 2.{m[1]}',text)
    P=[(heading,refs(text)) for heading,text in P]
    report=style_doc();report.add_paragraph('Part 2 Security Level Investing',style='Title')
    for i,(head,text) in enumerate(P):
        if i in [1,3,4,6]:report.add_paragraph(head.split(' ',1)[0]+' '+{'2A':'Data validation','2B':'Market proxy','2C':'Value portfolios','2D':'Tracking error and HML'}[head.split(' ',1)[0]],style='Heading 1')
        body=report.add_paragraph(text)
        if i==5:body.paragraph_format.keep_together=True
    report.save(DOC_DIR/'Part2_report.docx')
    md=['# Part 2 Security Level Investing','']
    for i,(head,text) in enumerate(P):md+=['<!-- P'+str(i+1)+' '+head+' -->',text,'']
    (OUT/'Part2_report.md').write_text('\n'.join(md))
    pd.DataFrame(LEDGER).to_csv(OUT/'report_claim_ledger.csv',index=False)
    # Compact tables assembled solely from existing results, no methodology changes.
    quality=pd.DataFrame([
      ['Full 2A scope','283,380 stock-months; 1,207 securities; 245 formation months'],
      ['Membership anomalies','Apr 1999: 182 departures / 166 arrivals; Dec 2018: -17.9%; Feb 2019: -38.8%'],
      ['Return fields versus absent rows','100% observed-row return coverage; 1,514 interior gaps in 205 securities'],
      ['Raw return range','-92.2% to +3,017.6%'],
      ['Main clipping rule','Monthly 1st/99th percentiles; linear interpolation; 5,723 changed rows; no deletions'],
      ['Alternative interpretation','0.5th/99.5th percentiles: 2,883 changed rows; main remains 1/99'],
      ['P/B midpoint and flag','Pb=0.50: 10,197 rows; Pb=0.50 & Bv<=0.03 flag: 5,720 (5,630 in 2B-2D window)']],columns=['Diagnostic','Saved finding'])
    cond=conditional[['Bv','n_stock_months','n_pb_eq_050','frequency_pb_eq_050_given_bv']].copy()
    cond.columns=['Bv score','Rows','Pb=0.50 rows','Conditional frequency'];cond['Bv score']=cond['Bv score'].map(lambda x:f'{x:.2f}');cond['Conditional frequency']=cond['Conditional frequency'].map(lambda x:f'{x:.2%}')
    rank=a5.pivot(index='test',columns='sample',values='median_rho').reindex(['P/B vs log(cap rank / book rank)','P/B vs cap rank within book-rank bins','P/B vs book rank within cap-rank bins']).reset_index()
    rank.columns=['Rank diagnostic','All observations','Without exact provisional flags']
    for col in rank.columns[1:]:rank[col]=rank[col].map(lambda x:f'{x:.4f}')
    perf=cp.loc[strategies+[market],['annualized_arithmetic_return','annualized_geometric_return','annualized_volatility','terminal_wealth']].copy()
    perf.loc['MSCI USA gross USD']=bp.loc['MSCI USA gross USD',['annualized_arithmetic_return','annualized_geometric_return','annualized_volatility','terminal_value_of_1']].to_numpy()
    for col in perf.columns[:3]:perf[col]=perf[col].map(lambda x:f'{x:.2%}')
    perf.terminal_wealth=perf.terminal_wealth.map(lambda x:f'${x:.3f}')
    perf=perf.reset_index();perf.columns=['Portfolio','Arithmetic annual return','Geometric CAGR','Annual volatility','Ending $1']
    concentration=final.loc[strategies+[market],['n_holdings','max_weight','effective_n','top10_weight']].copy().reset_index()
    concentration.columns=['Portfolio','Holdings','Maximum weight','Effective N','Top 10 weight'];concentration.Holdings=concentration.Holdings.astype(int)
    for col in ['Maximum weight','Top 10 weight']:concentration[col]=concentration[col].map(lambda x:f'{x:.4%}')
    concentration['Effective N']=concentration['Effective N'].map(lambda x:f'{x:.1f}')
    sens=robust[robust.strategy.isin(strategies)].pivot(index='strategy',columns=['universe','return_version'],values='annualized_arithmetic_return').reindex(strategies)
    sens=sens.reindex(columns=[('All observations (main)','Winsorized (main)'),('All observations (main)','Original (sensitivity)'),('Without provisional flags (sensitivity)','Winsorized (main)'),('Without provisional flags (sensitivity)','Original (sensitivity)')])
    sens.columns=['All clipped','All original','No flags clipped','No flags original'];sens=sens.map(lambda x:f'{x:.2%}').reset_index().rename(columns={'strategy':'Portfolio'})
    tracking=pd.DataFrame({'Portfolio':strategies,'Main TE':[D(s,'annualized_tracking_error') for s in strategies],'Original TE':[f'{dp.loc[("Original (sensitivity)",s),"annualized_tracking_error"]:.2%}' for s in strategies],'Main HML rho':[D(s,'HML_correlation','.4f') for s in strategies],'Original HML rho':[f'{dp.loc[("Original (sensitivity)",s),"HML_correlation"]:.4f}' for s in strategies]})
    cm=corr.copy();cm.columns=['S1 active','S2 active','S3 active','HML'];cm.index=['S1 active','S2 active','S3 active','HML'];cm=cm.map(lambda x:f'{x:.4f}').reset_index(names='Series')
    tables=[('T1','Data quality and cleaning',quality,[1.45,4.75],'Full 2A history; supplied Part 2 CSV. Formation Nov 1998-Mar 2019; all securities. Missing rows are diagnostic, not imputed returns.'),('T2a','Price to book midpoint conditional frequencies',cond,[1,1,1.5,2.7],'Same full history. Frequency denominator: all rows at the exact Bv score. Neither score identifies equity sign.'),('T2b','Ordinal sanity checks for price to book',rank,[3.2,1.2,1.8],'Saved 2A5 results, exact provisional flag. Proxy medians use 245 months; conditional medians use month-bin cells. Two singleton cells per sample in the book association are undefined and retained in the source audit.'),('T3','Main performance',perf,[1.8,1.1,1.1,1.1,1.1],'Investment Jan 1999-Dec 2018, 240 months. Supplied stock CSV and MSCI gross USD workbook. Main stock returns monthly 1/99 clipped; all flags retained; decimal inputs displayed as percentages. CAGR and arithmetic annualization differ.'),('T4','Final formation concentration',concentration,[1.8,.7,1.1,1.1,1.5],'Formation Nov 2018; supplied stock CSV. All observations, main weights. Holdings count positive weights; effective N = 1/sum(w squared).'),('T5','Arithmetic return sensitivities',sens,[1.8,1.1,1.1,1.1,1.1],'Jan 1999-Dec 2018. Same Value score and tie rule. Exclusion removes 5,630 flagged in-window stock-months only; full-universe benchmark remains fixed. No excluded observations are deleted from source files.'),('T6','Tracking error and HML sensitivity',tracking,[1.8,1.1,1.1,1.1,1.1],'Jan 1999-Dec 2018, 240 aligned observations. TE = sample standard deviation of Value-minus-2B market times sqrt(12). Both Value and market switch to original returns in sensitivity; HML unchanged.'),('T7','Complete active return correlation matrix',cm,[1.8,1.1,1.1,1.1,1.1],'Pearson correlations, main clipped returns. S1 EW top 20%; S2 Value score; S3 Value times cap. Saved 2B/2C and supplied monthly FF HML, Jan 1999-Dec 2018, all flags retained.')]
    appendix=style_doc(9);appendix.styles['Heading 1'].font.size=Pt(9);appendix.add_paragraph('Part 2 Selected Exhibits',style='Title')
    for idx,(tid,title,frame,widths,desc) in enumerate(tables):
        appendix.add_paragraph(refs(tid)+' '+title,style='Heading 1');addtable(appendix,frame,widths);note(appendix,'Source and scope: '+desc)
        frame.to_csv(OUT/(tid+'_display_table.csv'),index=False)
    figs=[('F1','Monthly stock counts','fig_2A_stock_counts'),('F2','Forward return distribution','fig_2A_return_distribution'),('F3','Return coverage and sample presence','fig_2A_return_coverage'),('F4','Price to book distribution and low book ranks','fig_2A_pb_distribution'),('F5','Rank sanity comparison for the exact provisional flag','fig_2A5_rank_sanity_flag_comparison'),('F6','Market proxy and MSCI USA cumulative wealth','fig_2B_cumulative_performance'),('F7','Value strategy cumulative wealth','fig_2C_cumulative_wealth'),('F8','Active return and HML correlations','fig_2D_correlation_main')]
    for idx,(fid,title,name) in enumerate(figs):
        if idx%2==0:appendix.add_page_break()
        appendix.add_paragraph(refs(fid)+' '+title,style='Heading 1')
        path=config.FIG_DIR/(name+'.png');im=Image.open(path)
        # Preserve saved print size and fonts rather than shrinking to fit.
        p=appendix.add_paragraph();p.paragraph_format.space_after=Pt(4);p.add_run().add_picture(str(path),width=Inches(im.width/300))
    appendix.save(DOC_DIR/'Part2_exhibits.docx')
    selection=[dict(id=refs(t[0]),type='table',title=t[1],source='Existing numerical tables; see display CSV and full inventory',scope=t[4]) for t in tables]
    selection += [dict(id=refs(fid),type='figure',title=title,source='output/figures/'+name+'.png',scope='Unchanged existing exhibit; source, dates and universe embedded') for fid,title,name in figs]
    pd.DataFrame(selection).to_csv(OUT/'exhibit_selection.csv',index=False)
    # Requirement mapping uses stable paragraph IDs in Markdown, not renderer-dependent page numbers.
    requirements=[
      ('2A-1','Counts, unusual months and universe consistency','P2','T1; F1'),('2A-2','Return distribution and 99% winsorization interpretation/cutoffs','P2','T1; F2; tab_2A_winsor_cutoffs.csv'),('2A-3','Return coverage; nulls versus disappearing/missing securities','P2','T1; F3; tab_2A_interior_gaps.csv'),('2A-4','P/B concentration and nonpositive-equity hypothesis','P3','T2a; F4; addendum_2A_pb_flag.md'),('2A-5','P/B accounting relation with percentile limitation','P3','T2b; F5; tab_2A5_conditional_cells.csv'),('2B-1','Score-cap weights, fully invested, formation versus earned month','P1, P4','T3; tab_2B_monthly_portfolio_audit.csv'),('2B-2','MSCI comparison and US-market/data-accuracy assessment','P4','T3; F6; tab_2B_msci_level_conversion.csv'),('2C-1','Three long-only Value methods and percentile interpretation','P5','T3; methodology_2C.md'),('2C-2','Average annualized return and appropriate mean','P1, P5','T3 (arithmetic and CAGR)'),('2C-3','Annualized volatility and economic comparison','P5','T3'),('2C-4','Final holdings count and highest single weight','P6','T4; tab_2C_final_holdings.csv'),('2D-1','Monthly active return differences and annualized TE','P7','T6; tab_2D_active_and_hml_main.csv'),('2D-2','Full 4x4 correlation and same 240 dates','P7','T7; F8; validation_2D.json'),('2D-3','Most HML-like strategy and HML construction explanation','P7','T6-T7; methodology_2D.md'),('QA-1','Original versus clipped return sensitivity','P6-P7','T5-T6; tab_2B_performance.csv'),('QA-2','Exact provisional-flag sensitivity without main exclusion','P3, P6','T2b; T5; holdings_2C.csv.gz'),('QA-3','Ties, endpoints and alternative clipping definition','P2, P5','T1; tab_2C_top20_boundary_tie_sensitivity.csv'),('QA-4','Shared MSCI USA and HML reference reconciliation','P4, P7','member1_reference_values.csv; Member 1 comparison pending'),('QA-5','Noncausal interpretation, data limits and unresolved choices','P3-P4, P6-P7','remaining_judgment_issues.md')]
    (OUT/'requirement_checklist.md').write_text('# Professor requirement coverage\n\nParagraph IDs P1-P7 correspond to the ordered prose paragraphs in Part2_report.md and the Word report. Detailed exhibits are outside the main-text budget.\n\n| ID | Requirement | Report paragraph | Exhibit or audit |\n|---|---|---|---|\n'+'\n'.join('| '+' | '.join(refs(value) for value in row)+' |' for row in requirements)+'\n')
    issues='''# Remaining judgment and coordination issues

1. Confirm instructor agreement that “99% winsorization” means monthly 1st/99th-percentile clipping (central 98%). This remains the established main method. The central-99% alternative is documented but no new strategy backtest is substituted.
2. Verify Pb=0.50 provenance and obtain equity signs if possible. The exact provisional flag is Pb=0.50 & Bv<=0.03, not all midpoint observations. The final 2A5 correlation without this flag is 0.8503; older exploratory exclusions of every midpoint produce different numbers (0.8544) and must not be confused. No numerical or methodological inconsistency was found in the final results.
3. Accept the fixed grid reversal v=1.01-Pb and deterministic stock-ID cutoff tie rule. Endpoint conventions and within-month reranking would change weights. The existing all-boundary-ties EW diagnostic changes main arithmetic annual return by -0.097 percentage points; do not silently replace the exact-k specification.
4. Historical extraction, 1999 membership breaks, survivorship, corporate events and point-in-time accounting cannot be certified from anonymized ranks. Positive correlations support aggregate consistency, not complete US coverage or individual-observation accuracy. Size/Value return attribution remains a hypothesis.
5. Original-return sensitivity materially changes Value performance, especially EW top 20%. Flag exclusion is a sensitivity, not correction of confirmed errors. Costs, turnover impacts and raw capitalization weights remain outside the established analysis.
6. Share member1_reference_values.csv with Member 1 and compare their independent loaders, including Jan 1995 MSCI USA 0.0278102318845 and HML 0.0257. Both series internally reconcile to the professor's workbook, but Member 1's independent calculations were not available and have NOT been declared matched. Use the supplied 2023-07 FF vintage; newer live French downloads may differ. HML wording must distinguish value weighting within component portfolios from equal small/big-leg averaging. No teammate messages are sent by this consolidation.
7. Group editor should fit this approximately 1.5-page prose into the final eight-page report, retain separately numbered exhibits, and add the actual member contribution on the shared cover. These shared materials are not edited here. No repository publication, commits or pushes before user review.
'''
    (OUT/'remaining_judgment_issues.md').write_text(issues)
    # Reconcile shared references directly to workbook values without invoking writers.
    sh=pd.read_excel(config.FILE_FF_MSCI,sheet_name=config.SHEET_FF3,header=None)
    months=sh.iloc[:,0].astype(str).str.fullmatch(r'\d{6}(\.0)?')
    ff=sh.loc[months,[0,3]].copy();ff.index=pd.to_datetime(ff[0].astype(str).str[:6],format='%Y%m').dt.to_period('M').astype(str)
    check('HML_shared_references_match_professor_workbook',np.allclose(pd.to_numeric(ff.loc[reference.index,3]).to_numpy()/100,reference.HML_decimal,atol=1e-12))
    sh=pd.read_excel(config.FILE_FF_MSCI,sheet_name=config.SHEET_MSCI,header=None)
    dates=pd.to_datetime(sh.iloc[:,0].astype(str),format='%b %d, %Y',errors='coerce');rows=dates.notna()
    ix=dates[rows].dt.to_period('M').astype(str)
    levels=pd.Series(pd.to_numeric(sh.loc[rows,1].astype(str).str.replace(',','',regex=False),errors='raise').to_numpy(),index=ix)
    check('MSCI_shared_references_match_professor_workbook',np.allclose(levels.loc[reference.index],reference.current_level,atol=1e-9) and np.allclose(levels.loc[(pd.PeriodIndex(reference.index,freq='M')-1).astype(str)],reference.prior_level,atol=1e-9))
    checksfiles={}
    for path in sorted((ROOT/'part2').glob('validation_*.json')):
        record=json.loads(path.read_text());checksfiles[path.name]=len(record.get('checks',{}));check('prior_checks_pass_'+path.stem,all(x=='PASS' for x in record.get('checks',{}).values()))
    (OUT/'consolidated_qa.json').write_text(json.dumps({'methodology_changed':False,'analysis_scripts_rerun':False,'checks':CHECKS,'prior_check_counts':checksfiles,'member1_independent_comparison':'pending; internal source references reconciled','report_word_count':sum(len(p.split()) for _,p in P)},indent=2))
    pd.DataFrame(LEDGER).to_csv(OUT/'report_claim_ledger.csv',index=False)
    # Browser-readable companion with the same text and the selected existing exhibits.
    page=['<!doctype html><html lang="en"><meta charset="utf-8"><title>FE5108 Part 2 consolidated report</title><style>body{font:14px Arial;line-height:1.5;max-width:1000px;margin:36px auto}h1{font-size:23px}h2{font-size:18px;margin-top:24px}table{border-collapse:collapse;font-size:12px;width:100%;margin:12px 0}th,td{padding:8px;border-bottom:1px solid #ddd}th{background:#e8eef2}img{width:auto;max-width:100%;height:auto}p.note{font-size:11px;color:#555}</style><h1>Part 2 Security Level Investing</h1>']
    for i,(head,text) in enumerate(P):page.append(f'<p id="P{i+1}">'+html.escape(text)+'</p>')
    page.append('<h1>Selected exhibits outside main text</h1>')
    for tid,title,frame,widths,desc in tables:page+=['<h2>'+refs(tid)+' '+title+'</h2>',frame.to_html(index=False),'<p class="note">'+html.escape(desc)+'</p>']
    for fid,title,name in figs:page+=['<h2>'+refs(fid)+' '+title+'</h2>',f'<img src="../../../output/figures/{name}.png" alt="{title}">']
    page+=['<h2>Open judgments and coordination</h2><pre style="white-space:pre-wrap">'+html.escape(issues)+'</pre></html>']
    (OUT/'Part2_review.html').write_text('\n'.join(page))
    print('Consolidated checks:',len(CHECKS),'PASS; prose words:',sum(len(p.split()) for _,p in P))

if __name__=='__main__':main()
