"""Build the Part 3 report and exhibit appendix (Word), mirroring part2/python_source/build_consolidated.py.

Run AFTER question_3A ... question_3E notebooks:
    python part3/python_source/build_part3_docs.py

Inputs : output/tables/tab_3*.csv, output/figures/fig_3*.png, weights_3A.pkl and the raw data
         (a few statistics quoted in the text are recomputed here and checked against the saved tables).
Outputs: part3/documents/Part3_report.docx, part3/documents/Part3_exhibits.docx
Rerunning overwrites both Word files - keep manual edits in a separate copy.
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import part3_common as p3  # noqa: E402
from part3_common import config, utils, EN  # noqa: E402

ROOT = p3.ROOT
TAB = config.TAB_DIR
DOC_DIR = ROOT / "part3" / "documents"
TEXT_WIDTH_IN = 8.2677 - 2.0          # A4 minus 1-inch side margins (same page set-up as Part 2)
CHECKS = {}


def check(name, cond):
    if not bool(cond):
        raise AssertionError(name)
    CHECKS[name] = "PASS"


def read(name, **kw):
    return pd.read_csv(TAB / name, **kw)


def pct(x, d=1):
    return f"{x * 100:.{d}f}%"


# --------------------------------------------------------------------------- Word helpers (same look as Part 2)
def style_doc(fontsize=11, title="FE5108 Part 3 Portfolio Construction"):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.2677), Inches(11.6929)
    s.top_margin = s.bottom_margin = Inches(.85)
    s.left_margin = s.right_margin = Inches(1)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        st = d.styles[name]
        st.font.name = "Arial"
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    d.styles["Normal"].font.size = Pt(fontsize)
    d.styles["Normal"].paragraph_format.line_spacing = 1.08
    d.styles["Normal"].paragraph_format.space_after = Pt(6)
    d.styles["Title"].font.size = Pt(15)
    d.styles["Title"].paragraph_format.space_after = Pt(9)
    d.styles["Heading 1"].font.size = Pt(11)
    d.styles["Heading 1"].paragraph_format.space_before = Pt(8)
    d.styles["Heading 1"].paragraph_format.space_after = Pt(4)
    for border in list(d.styles.element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
    d.core_properties.title = title
    d.core_properties.author = "FE5108 Part 3"
    return d


def note(d, text):
    p = d.add_paragraph(text)
    p.paragraph_format.space_after = Pt(5)
    for r in p.runs:
        r.font.size = Pt(6.5)
        r.font.color.rgb = RGBColor.from_string("555555")


def addtable(d, frame, widths, fills=None, bold_rows=(), fontsize=8, bold_cells=()):
    """Part 2 table style; `fills` = {(row, col): 'RRGGBB'} for data cells (row 0 = first data row)."""
    fills = fills or {}
    t = d.add_table(rows=1, cols=len(frame.columns))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths):
        col.width = Inches(w)
    for i, (cell, name) in enumerate(zip(t.rows[0].cells, frame.columns)):
        cell.text = str(name)
        cell.width = Inches(widths[i])
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in frame.itertuples(index=False, name=None):
        for i, (cell, value) in enumerate(zip(t.add_row().cells, row)):
            cell.text = str(value)
            cell.width = Inches(widths[i])
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for side in ["top", "bottom", "left", "right"]:
                e = OxmlElement("w:" + side)
                e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "D9D9D9")
                borders.append(e)
            tcPr.append(borders)
            margins = OxmlElement("w:tcMar")
            for side in ["top", "bottom", "left", "right"]:
                e = OxmlElement("w:" + side)
                e.set(qn("w:w"), "55"); e.set(qn("w:type"), "dxa")
                margins.append(e)
            tcPr.append(margins)
            fill = "E8EEF2" if ri == 0 else fills.get((ri - 1, ci))
            if fill:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:val"), "clear"); shade.set(qn("w:color"), "auto"); shade.set(qn("w:fill"), fill)
                tcPr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = True
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 or str(frame.columns[ci]).startswith("Industry") \
                    else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(fontsize)
                    run.bold = ri == 0 or (ri - 1) in bold_rows or (ri - 1, ci) in bold_cells
    spacer = d.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(2)
    spacer.paragraph_format.keep_with_next = True
    return t


# --------------------------------------------------------------------------- statistics
def stats():
    S = {}
    rets = p3.load_industry_returns()
    W = p3.load_weights_3A()
    seed, oos = utils.select_period(rets, "3_seed"), utils.select_period(rets, "3C")
    eig = np.linalg.eigvalsh(seed.cov().to_numpy())
    S["min_eig"], S["cond"] = eig.min(), eig.max() / eig.min()
    ws = W["mv_short"]
    S["long_sum"], S["short_sum"] = ws.clip(lower=0).sum(), ws.clip(upper=0).sum()
    S["n_long"], S["n_short"] = int((ws > 0).sum()), int((ws < 0).sum())
    L, Sh = ws.clip(lower=0), (-ws).clip(lower=0)
    for name, df in [("is", seed), ("oos", oos)]:
        S[f"leg_corr_{name}"] = (df @ (L / L.sum())).corr(df @ (Sh / Sh.sum()))
    ac = lambda df: df.corr().to_numpy()[np.triu_indices(49, 1)].mean()
    S["avg_corr_is"], S["avg_corr_oos"] = ac(seed), ac(oos)
    S["med_vol_is"] = seed.apply(utils.annualized_volatility).median()
    S["med_vol_oos"] = oos.apply(utils.annualized_volatility).median()
    R = p3.portfolio_returns(oos, W)
    roll = (R.rolling(36).std() * np.sqrt(12)).dropna()
    S["roll_share"] = (roll["mv_long_only"] < roll[["mv_short", "inverse_vol"]].min(axis=1)).mean()
    S["n_roll"] = len(roll)
    # 3D
    seed2, test = utils.select_period(rets, "3D_seed"), utils.select_period(rets, "3D_oos")
    w_old = W["mv_long_only"]
    w_new = p3.min_var_long_only(seed2)
    S["w_new"], S["turnover"] = w_new, 0.5 * (w_new - w_old).abs().sum()
    pf = lambda w, df: df[w.index] @ w
    S["old_0615"] = utils.annualized_volatility(pf(w_old, seed2))
    S["new_0615"] = utils.annualized_volatility(pf(w_new, seed2))
    cov = lambda r: (1 + r.loc["2020-02":"2020-03"]).prod() - 1
    S["covid_old"], S["covid_new"], S["covid_ew"] = cov(pf(w_old, test)), cov(pf(w_new, test)), cov(test.mean(axis=1))
    return S, W


def main():
    DOC_DIR.mkdir(parents=True, exist_ok=True)
    S, W = stats()

    # ---- saved tables
    w3a = read("tab_3A_weights.csv", index_col="Code")
    oosv = read("tab_3C_oos_vol.csv", index_col="Portfolio")
    sub = read("tab_3C_subperiod_vol.csv", index_col="Sub-period")
    w3d = read("tab_3D_weights.csv", index_col="Code")
    v3d = read("tab_3D_reopt_vol.csv", index_col="Period (annualized volatility)")
    perf = read("tab_3E_performance.csv", index_col=0)
    contrib = read("tab_3E_crisis_contrib.csv", index_col=["Crisis", "Code"])
    top = {k: read(f"tab_3E_top10_{k.lower()}.csv", index_col="Code") for k in ["GFC", "COVID"]}

    cols = list(p3.COLS.values())
    wts = w3a.drop(index="Sum")[cols]
    check("3A_weights_match_pickle", np.allclose(wts.to_numpy(), W.rename(columns=p3.COLS)[cols].loc[wts.index].to_numpy(), atol=1e-6))
    check("3A_weights_sum_to_one", np.allclose(wts.sum(), 1, atol=1e-5))
    check("3D_new_weights_match_table", np.allclose(S["w_new"].loc[w3d.index], w3d["New weights (as of Dec 2015)"], atol=1e-6))
    ins, oo = oosv.iloc[:, 0], oosv.iloc[:, 1]
    sh, lo, iv = cols
    ew = "Equal weight (49 industries)"
    check("3C_ordering_in_sample", ins[sh] < ins[lo] < ins[iv])
    check("3C_long_only_lowest_oos", oo[lo] < oo[sh] < oo[iv])
    check("3C_long_only_lowest_every_subperiod", (sub["Lowest-vol portfolio"] == lo).all())
    old_c, new_c = "Old (Dec 2004 weights)", "New (Dec 2015 weights)"
    full = v3d.index[0]

    gfc_l, cov_l = "Cumulative return, GFC (Sep 2008 - Feb 2009)", "Cumulative return, COVID (Feb - Mar 2020)"
    pm, mm = perf["Long-only min-var"], perf["MSCI USA"]
    re_g, re_c = contrib.loc[("GFC", "RlEst")], contrib.loc[("COVID", "RlEst")]
    ret_col, ctr_col, rank_col = "Industry crisis return", "Contribution (≈ weight × return)", "Rank of 49 (1 = best)"
    wl = wts[lo]
    held = wl[wl > 0].sort_values(ascending=False)
    wn = S["w_new"][S["w_new"] > 0].sort_values(ascending=False)
    ut_c = contrib.loc[("COVID", "Util")]

    # ---- report text (English, own words); references T = Table 3.x, F = Figure 3.x
    P = []
    P.append(("3A", "We use value-weighted monthly total returns of the 49 Fama-French industry portfolios, January 1995 to "
              "July 2023; the -99.99/-999 missing-value codes do not occur in this window. Weights are set once at the close of "
              "31 December 2004 from the sample covariance matrix of January 1995-December 2004 (120 months) and then held "
              "constant, rebalanced monthly to target. Volatility is the sample standard deviation of monthly returns times √12."))
    P.append(("3A", "The long-only minimum-variance portfolio solves: minimise w′Σw subject to Σᵢwᵢ = 1 and wᵢ ≥ 0 for all 49 "
              f"industries. Σ is positive definite (condition number {S['cond']:,.0f}), so this convex quadratic programme has a "
              "unique solution. We solve it with cvxpy (CLARABEL interior-point solver), confirm it with SciPy SLSQP (largest weight "
              "difference below 10⁻⁷) and verify the KKT conditions: every held industry has the same marginal contribution to "
              "portfolio variance and every excluded industry a higher one. The minimum-variance portfolio with shorting uses the "
              "closed form w = Σ⁻¹1/(1′Σ⁻¹1); the inverse-volatility portfolio sets wᵢ proportional to 1/σᵢ. Code is in the appendix."))
    P.append(("3A", f"Allowing shorts produces an extreme portfolio: {S['n_long']} long and {S['n_short']} short positions, "
              f"{pct(S['long_sum'], 0)} long and {pct(S['short_sum'], 0)} short, led by {EN[ws_max(W)]} ({pct(W['mv_short'].max())}) "
              f"and {EN[W['mv_short'].idxmin()]} ({pct(W['mv_short'].min())}). The long-only solution is a corner solution with "
              f"only {len(held)} industries, led by {EN[held.index[0]]} ({pct(held.iloc[0])}), {EN[held.index[1]]} ({pct(held.iloc[1])}) "
              f"and {EN[held.index[2]]} ({pct(held.iloc[2])}); the other {49 - len(held)} weights are zero. Inverse volatility holds "
              f"all 49 industries at {pct(wts[iv].min())}-{pct(wts[iv].max())}. Utilities and Real estate are among the five largest "
              "weights under all three methods. See T1-T2 and F1."))
    P.append(("3B", f"In-sample annualized volatility is {pct(ins[sh])} with shorts, {pct(ins[lo])} long-only and {pct(ins[iv])} "
              f"for inverse volatility, against {pct(ins[ew])} for an equally weighted portfolio. The ordering is as expected and in "
              "fact guaranteed: because Σ was estimated on these same returns, realized in-sample volatility equals the optimizer's "
              "objective, and each restriction (no shorting) or simplification (ignoring correlations) can only raise the attainable "
              "minimum. The figure therefore rewards fitting the sample, including its estimation error: 120 observations are few "
              "for a 49×49 covariance matrix, and the unconstrained optimizer exploits that error most aggressively. See T3 and F2."))
    P.append(("3C", f"Out-of-sample, January 2005-July 2023, the long-only portfolio is lowest at {pct(oo[lo])}, followed by "
              f"shorts allowed at {pct(oo[sh])} and inverse volatility at {pct(oo[iv])} (equal weight {pct(oo[ew])}). The deterioration "
              f"follows how freely each method uses the estimated covariances: volatility rises ×{oo[sh] / ins[sh]:.2f}, "
              f"×{oo[lo] / ins[lo]:.2f} and ×{oo[iv] / ins[iv]:.2f}, versus ×{oo[ew] / ins[ew]:.2f} for equal weight. This is the "
              f"error-maximization effect discussed in class. With {S['long_sum'] - S['short_sum']:.1f}× gross leverage, a small "
              f"drop in the correlation between the long and short legs ({S['leg_corr_is']:.3f} to {S['leg_corr_oos']:.3f}) more "
              "than doubles risk. The long-only constraint acts as shrinkage of the covariance estimate (Jagannathan and Ma, 2003). "
              f"Average pairwise industry correlation also rose from {S['avg_corr_is']:.2f} to {S['avg_corr_oos']:.2f}, while median "
              f"industry volatility was unchanged ({pct(S['med_vol_is'])} vs {pct(S['med_vol_oos'])}), weakening diversification "
              f"for every portfolio. The long-only portfolio is lowest in all four sub-periods and in {pct(S['roll_share'], 0)} of "
              f"{S['n_roll']} rolling 36-month windows; in 2008-2009 the unconstrained portfolio ({pct(sub.loc['2008–2009 (GFC)', sh])}) "
              f"was even riskier than equal weight ({pct(sub.loc['2008–2009 (GFC)', ew])}). See T3-T4 and F3-F4."))
    P.append(("3D", f"Re-optimizing at end-2015 on January 2006-December 2015 returns gives a portfolio of {len(wn)} industries: "
              + ", ".join(f"{EN[c]} {pct(v)}" for c, v in wn.items())
              + f". Turnover from the 2004 weights is {pct(S['turnover'])} and Real estate is dropped entirely. Over January "
              f"2016-July 2023 its volatility is {pct(v3d.loc[full, new_c])} versus {pct(v3d.loc[full, old_c])} for the 2004 "
              f"portfolio ({pct(v3d.loc[full, ew])} equal weight), a {pct(1 - v3d.loc[full, new_c] / v3d.loc[full, old_c], 0)} "
              "relative reduction. This is expected because covariances are not stable: over 2006-2015 the stale 2004 weights had "
              f"{pct(S['old_0615'])} volatility against {pct(S['new_0615'])} for the re-optimized ones. The gain is concentrated in the "
              f"COVID period, however: before 2020 the two are close ({pct(v3d.iloc[1][new_c])} vs {pct(v3d.iloc[1][old_c])}), whereas "
              f"in February-March 2020 the new portfolio lost {pct(-S['covid_new'])} against {pct(-S['covid_old'])}. Re-optimization "
              f"is no free lunch: estimation error remains (out-of-sample volatility {pct(v3d.loc[full, new_c] / S['new_0615'] - 1, 0)} "
              f"above in-sample), {pct(wn.iloc[0])} sits in one industry, and transaction costs are ignored. See T5-T6 and F5-F6."))
    P.append(("3E", f"$1 invested on 31 December 2004 grew to ${pm.iloc[0]:.2f}, versus ${mm.iloc[0]:.2f} in the MSCI USA index, "
              f"with lower volatility ({pct(pm.iloc[2])} vs {pct(mm.iloc[2])}) and a smaller maximum drawdown ({pct(pm.iloc[3])} vs "
              f"{pct(mm.iloc[3])}). The defensive strategy did not, however, deliver as expected in the two crises: it lost "
              f"{pct(-pm[gfc_l])} in the GFC against {pct(-mm[gfc_l])} for the market, and {pct(-pm[cov_l])} in February-March 2020, "
              f"more than the market's {pct(-mm[cov_l])}. The main reason is the {pct(re_g['Weight'])} weight in Real estate, a "
              f"low-volatility industry in 1995-2004 but among the five worst in both crises ({pct(re_g[ret_col], 0)} and "
              f"{pct(re_c[ret_col], 0)}), costing roughly {pct(-re_g[ctr_col])} and {pct(-re_c[ctr_col])} of portfolio return. "
              f"Utilities, the largest weight, fell {pct(-ut_c[ret_col], 0)} in the COVID sell-off. See T7-T8 and F7."))
    g, c = top["GFC"], top["COVID"]
    vr = "Vol rank 1995–2004"
    P.append(("3E", "The industries that held up best in the GFC were mostly traditional defensives (food, beverages, "
              "pharmaceuticals, tobacco, utilities, retail), plus Precious metals, the best performer "
              f"({pct(g.loc['Gold', 'Crisis return'], 0)}) although it ranked {g.loc['Gold', vr]} of 49 by volatility: a safe haven is "
              "not a low-volatility asset. The COVID list adds stay-at-home beneficiaries such as Computer software, Retail, Shipping "
              f"containers and Measuring & control equipment. Software ranked {c.loc['Softw', vr]} and Measuring & control equipment "
              f"{c.loc['LabEq', vr]} of 49 by 1995-2004 volatility, so a minimum-variance portfolio excludes them by design; only "
              f"{int((g['Min-var weight'] > 0).sum())} of the ten best industries in each crisis were held. Historical low volatility is "
              "therefore no guarantee of crisis protection: which industries are defensive depends on the nature of the shock, and "
              "both covariances and the set of low-volatility industries change over time. See T9 and F8."))

    def refs(text):
        import re
        text = re.sub(r"\bT(\d)-T(\d)\b", lambda m: f"Tables 3.{m[1]}-3.{m[2]}", text)
        text = re.sub(r"\bF(\d)-F(\d)\b", lambda m: f"Figures 3.{m[1]}-3.{m[2]}", text)
        text = re.sub(r"\bT(\d)\b", lambda m: f"Table 3.{m[1]}", text)
        return re.sub(r"\bF(\d)\b", lambda m: f"Figure 3.{m[1]}", text)

    heads = {"3A": "3A Minimum-variance and inverse-volatility portfolios", "3B": "3B In-sample effectiveness",
             "3C": "3C Out-of-sample effectiveness", "3D": "3D The value of re-optimization",
             "3E": "3E Performance expectations"}
    report = style_doc()
    report.add_paragraph("Part 3 Portfolio Construction", style="Title")
    last = None
    for q, text in P:
        if q != last:
            report.add_paragraph(heads[q], style="Heading 1")
            last = q
        report.add_paragraph(refs(text))
    ref = report.add_paragraph("Reference: Jagannathan, R. and Ma, T. (2003). Risk reduction in large portfolios: Why imposing "
                               "the wrong constraints helps. Journal of Finance, 58(4).")
    for r in ref.runs:
        r.font.size = Pt(8)
    report.save(DOC_DIR / "Part3_report.docx")
    words = sum(len(refs(t).split()) for _, t in P)

    # ---- exhibit tables
    tables = []
    # T1 weights, two side-by-side halves with green / red marks
    fills = {}
    def marks(col):
        s = wts[col]
        top5 = set(s.nlargest(5).index)
        bot = set(s.index[s == 0]) if col == lo else set(s.nsmallest(5).index)
        return top5, bot
    mk = {col: marks(col) for col in cols}
    short_hdr = ["Min-var (shorts)", "Min-var (long-only)", "Inverse vol"]
    left, right = wts.index[:25], wts.index[25:]
    rows = []
    for i in range(25):
        row = []
        for half, offset in [(left, 0), (right, 4)]:
            if i < len(half):
                code = half[i]
                row += [EN[code]] + [pct(wts.at[code, col]) for col in cols]
                for j, col in enumerate(cols):
                    if code in mk[col][0]:
                        fills[(i, offset + 1 + j)] = p3.light_hex(config.COLOR_POS)
                    elif code in mk[col][1]:
                        fills[(i, offset + 1 + j)] = p3.light_hex(config.COLOR_NEG)
            else:
                row += ["Sum of weights"] + [pct(wts[col].sum()) for col in cols]
        rows.append(row)
    t1 = pd.DataFrame(rows, columns=["Industry"] + short_hdr + ["Industry "] + [h + " " for h in short_hdr])
    n_zero = int((wts[lo] == 0).sum())
    tables.append(("T1", "Optimal industry weights as of 31 December 2004", t1, [1.25, .58, .62, .58] * 2, fills, {(24, k) for k in range(4, 8)}, 7.5,
                   "49 Fama-French industries, value-weighted total returns; covariance estimated Jan 1995 - Dec 2004. Fully invested. "
                   f"Green = 5 largest weights per portfolio; red = 5 smallest (long-only: all {n_zero} zero-weight industries). "
                   "Industries in Fama-French order."))
    # T2 characteristics
    def prof(s):
        return [pct(s.sum()), int((s.abs() > 1e-12).sum()), int((s < 0).sum()), pct(s.max()), pct(s.min()),
                pct(s.nlargest(5).sum()), f"{1 / (s ** 2).sum():.1f}"]
    t2 = pd.DataFrame({h: prof(wts[c]) for h, c in zip(cols, cols)},
                      index=["Sum of weights", "Non-zero positions", "Short positions", "Largest weight", "Smallest weight",
                             "Top-5 weight", "Effective N = 1/Σw²"]).reset_index(names="Characteristic")
    tables.append(("T2", "Portfolio characteristics as of 31 December 2004", t2, [1.9, 1.45, 1.45, 1.45], {}, (), 8,
                   "Weights from Table 3.1. Effective N equals 49 for equal weights; it is not meaningful for the portfolio with short positions."))
    # T3 in-sample vs out-of-sample
    t3 = pd.DataFrame({"Portfolio": oosv.index,
                       "In-sample vol, Jan 1995 - Dec 2004": oosv.iloc[:, 0].map(pct),
                       "Out-of-sample vol, Jan 2005 - Jul 2023": oosv.iloc[:, 1].map(pct),
                       "Ratio OOS / in-sample": oosv.iloc[:, 2].map(lambda x: f"×{x:.2f}")})
    tables.append(("T3", "In-sample versus out-of-sample annualized volatility", t3, [2.5, 1.25, 1.35, 1.1], {}, (), 8,
                   "Weights from Table 3.1 held constant and rebalanced monthly. Reference rows (equal weight; lowest-volatility "
                   "industry in 1995-2004, Food products) use the same data and only information available at Dec 2004."))
    # T4 sub-periods
    t4 = sub.copy()
    for col in t4.columns[:-1]:
        t4[col] = t4[col].map(pct)
    t4 = t4.drop(columns="Lowest-vol portfolio").reset_index().rename(columns={"Sub-period": "Out-of-sample sub-period"})
    t4.columns = ["Out-of-sample sub-period", "Min-var (shorts)", "Min-var (long-only)", "Inverse vol", "Equal weight"]
    tables.append(("T4", "Out-of-sample volatility by sub-period", t4, [1.9, 1.1, 1.1, 1.1, 1.1], {}, (), 8,
                   "Annualized volatility of monthly returns within each sub-period; 3A weights, no re-estimation. "
                   "The long-only portfolio is the lowest of the three in every sub-period."))
    # T5 old vs new weights
    t5 = pd.DataFrame({"Industry": [EN[c] for c in w3d.index],
                       "Weights as of Dec 2004": w3d.iloc[:, 1].map(pct),
                       "Weights as of Dec 2015": w3d.iloc[:, 2].map(pct),
                       "Change": w3d.iloc[:, 3].map(lambda x: f"{x * 100:+.1f} pp")})
    t5.loc[len(t5)] = ["Turnover (½Σ|Δw|)", "", "", pct(S["turnover"])]
    tables.append(("T5", "Long-only minimum-variance weights before and after re-optimization", t5, [2.2, 1.3, 1.3, 1.3], {}, (len(t5) - 1,), 8,
                   "Dec 2004: covariance Jan 1995 - Dec 2004 (Table 3.1). Dec 2015: same method, covariance Jan 2006 - Dec 2015. "
                   "Industries with a non-zero weight in at least one portfolio."))
    # T6 re-optimization volatility
    t6 = v3d.copy()
    t6 = pd.DataFrame({"Period": t6.index, "Dec 2004 weights": t6[old_c].map(pct), "Dec 2015 weights": t6[new_c].map(pct),
                       "Equal weight": t6.iloc[:, 2].map(pct), "New − old": t6.iloc[:, 3].map(lambda x: f"{x * 100:+.1f} pp")})
    t6.loc[len(t6)] = ["Cumulative return, Feb - Mar 2020", pct(S["covid_old"]), pct(S["covid_new"]), pct(S["covid_ew"]),
                       f"{(S['covid_new'] - S['covid_old']) * 100:+.1f} pp"]
    tables.append(("T6", "Annualized volatility, January 2016 - July 2023", t6, [2.2, 1.0, 1.0, 1.0, 1.0], {}, (), 8,
                   "Both portfolios out-of-sample; weights held constant and rebalanced monthly. Last row: compounded return over "
                   "Feb - Mar 2020 (two months are too few for a volatility estimate)."))
    # T7 performance
    t7 = perf.copy()
    t7 = pd.DataFrame({"Jan 2005 - Jul 2023": t7.index,
                       "Long-only min-var": [f"${v:.2f}" if i == 0 else pct(v) for i, v in enumerate(t7.iloc[:, 0])],
                       "MSCI USA": [f"${v:.2f}" if i == 0 else pct(v) for i, v in enumerate(t7.iloc[:, 1])]})
    tables.append(("T7", "Performance of the Dec 2004 long-only minimum-variance portfolio and MSCI USA", t7, [3.3, 1.45, 1.45], {}, (), 8,
                   "Fama-French 49 industries (all CRSP stocks, value-weighted) and MSCI USA Standard (large + mid cap), gross, USD. "
                   "Annualized return is geometric (compound growth). Crisis returns compounded over the stated months."))
    # T8 crisis contributions
    cg, cc = contrib.loc["GFC"], contrib.loc["COVID"]
    order = cg["Weight"].sort_values(ascending=False).index
    t8 = pd.DataFrame({"Industry held": [EN[c] for c in order], "Weight": cg.loc[order, "Weight"].map(pct),
                       "GFC return": cg.loc[order, ret_col].map(lambda x: pct(x, 0)),
                       "GFC contribution": cg.loc[order, ctr_col].map(pct),
                       "COVID return": cc.loc[order, ret_col].map(lambda x: pct(x, 0)),
                       "COVID contribution": cc.loc[order, ctr_col].map(pct)})
    re_row = list(order).index("RlEst")
    tables.append(("T8", "Crisis returns of the portfolio's holdings", t8, [1.7, .75, .85, .95, .85, .95],
                   {(re_row, k): p3.light_hex(config.COLOR_NEG) for k in range(6)}, (), 8,
                   "GFC: Sep 2008 - Feb 2009; COVID: Feb - Mar 2020. Contribution ≈ weight × industry crisis return (approximate, as "
                   "the portfolio is rebalanced monthly). Real estate (shaded) was among the five worst industries in both crises."))
    # T9 top 10
    pre_cols = {k: [c for c in top[k].columns if c.startswith("Vol rank, 10y")][0] for k in top}
    rows = []
    for i in range(10):
        r = []
        for k in ["GFC", "COVID"]:
            t = top[k]
            code = t.index[i]
            r += [EN[code], pct(t["Crisis return"].iloc[i], 0), f"{t[vr].iloc[i]} / {t[pre_cols[k]].iloc[i]}",
                  pct(t["Min-var weight"].iloc[i]) if t["Min-var weight"].iloc[i] > 0 else "–"]
        rows.append(r)
    t9 = pd.DataFrame(rows, columns=["Industry (GFC)", "Return", "Vol rank", "Weight",
                                     "Industry (COVID)", "Return ", "Vol rank ", "Weight "])
    tables.append(("T9", "Ten best-performing industries in each crisis", t9, [1.25, .55, .6, .55] * 2, {}, (), 7.5,
                   "Vol rank = rank of annualized volatility among 49 industries (1 = lowest): Jan 1995 - Dec 2004 / ten years before "
                   "the crisis. Weight = Dec 2004 long-only minimum-variance weight (– = not held)."))

    appendix = style_doc(9, "FE5108 Part 3 Selected Exhibits")
    appendix.styles["Heading 1"].font.size = Pt(9)
    appendix.add_paragraph("Part 3 Selected Exhibits", style="Title")
    for tid, title, frame, widths, fills_, bold, fs, desc in tables:
        appendix.add_paragraph(refs(tid) + " " + title, style="Heading 1")
        if isinstance(bold, set):
            addtable(appendix, frame, widths, fills_, (), fs, bold_cells=bold)
        else:
            addtable(appendix, frame, widths, fills_, bold, fs)
        note(appendix, "Source and scope: Kenneth French Data Library; MSCI. " + desc)
    figs = [("F1", "Sorted industry weights of the three portfolios (3A)", "fig_3A_weight_profile"),
            ("F2", "In-sample volatility, Jan 1995 - Dec 2004 (3B)", "fig_3B_insample_vol"),
            ("F3", "In-sample versus out-of-sample volatility (3B-3C)", "fig_3C_oos_vol"),
            ("F4", "Rolling 36-month out-of-sample volatility (3C)", "fig_3C_rolling_vol"),
            ("F5", "Value of re-optimization: volatility, Jan 2016 - Jul 2023 (3D)", "fig_3D_reopt_vol"),
            ("F6", "Weights before and after re-optimization (3D)", "fig_3D_weights"),
            ("F7", "Growth of $1 invested on 31 Dec 2004, log2 scale (3E)", "fig_3E_cumulative"),
            ("F8", "Ten best-performing industries in each crisis (3E)", "fig_3E_top10")]
    scaled = []
    for idx, (fid, title, name) in enumerate(figs):
        if idx % 2 == 0:
            appendix.add_page_break()
        appendix.add_paragraph(refs(fid) + " " + title, style="Heading 1")
        path = config.FIG_DIR / (name + ".png")
        im = Image.open(path)
        width = im.width / 300                       # saved print size (300 dpi), as in Part 2
        if width > TEXT_WIDTH_IN:                    # only shrink figures wider than the text block
            scaled.append((name, TEXT_WIDTH_IN / width))
            width = TEXT_WIDTH_IN
        p = appendix.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(str(path), width=Inches(width))
    appendix.save(DOC_DIR / "Part3_exhibits.docx")

    print(f"Checks passed: {len(CHECKS)}  ({', '.join(CHECKS)})")
    print(f"Report: {DOC_DIR / 'Part3_report.docx'}  ({words} words of body text)")
    print(f"Exhibits: {DOC_DIR / 'Part3_exhibits.docx'}  ({len(tables)} tables, {len(figs)} figures)")
    for name, f in scaled:
        print(f"  note: {name}.png is wider than the text block, scaled to {f:.0%}")


def ws_max(W):
    return W["mv_short"].idxmax()


if __name__ == "__main__":
    main()
